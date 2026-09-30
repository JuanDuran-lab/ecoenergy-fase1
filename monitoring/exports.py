"""
Exportación de lecturas a Excel (.xlsx) con openpyxl.

Recibe un QuerySet ya filtrado por permisos, organización y filtros del
listado; construye el libro en memoria y lo devuelve para su descarga.
"""

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="15803D")

READING_COLUMNS = [
    ("ID", 8, None),
    ("Fecha y hora", 18, "DD-MM-YYYY HH:MM"),
    ("Organización", 20, None),
    ("Zona", 24, None),
    ("Dispositivo", 26, None),
    ("N° de serie", 16, None),
    ("Categoría", 16, None),
    ("Consumo (kWh)", 15, "#,##0.00"),
    ("Nominal (kWh)", 15, "#,##0.00"),
    ("Diferencia (kWh)", 16, "#,##0.00"),
    ("Sobre nominal", 14, None),
    ("Observaciones", 30, None),
]


def _local_naive(value):
    return timezone.localtime(value).replace(tzinfo=None)


def build_readings_workbook(queryset, *, user, filters_description=""):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lecturas"

    sheet.append([header for header, _, _ in READING_COLUMNS])
    for index, (_, width, _) in enumerate(READING_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width

    total_rows = 0
    for reading in queryset.iterator(chunk_size=500):
        device = reading.device
        difference = reading.consumption_kwh - device.nominal_consumption_kwh
        sheet.append(
            [
                reading.pk,
                _local_naive(reading.reading_at),
                device.zone.organization.name,
                device.zone.name,
                device.name,
                device.serial_number,
                device.category.name,
                float(reading.consumption_kwh),
                float(device.nominal_consumption_kwh),
                float(difference),
                "Sí" if difference > 0 else "No",
                reading.notes,
            ]
        )
        total_rows += 1

    for index, (_, _, number_format) in enumerate(READING_COLUMNS, start=1):
        if number_format:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = number_format

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    info = workbook.create_sheet("Información")
    organization = getattr(getattr(user, "profile", None), "organization", None)
    for label, value in [
        ("Reporte", "Lecturas de consumo - EcoEnergy"),
        ("Generado el", _local_naive(timezone.now())),
        ("Generado por", user.get_username()),
        ("Alcance", organization.name if organization else "Todas las organizaciones"),
        ("Filtros aplicados", filters_description or "Ninguno"),
        ("Total de registros", total_rows),
    ]:
        info.append([label, value])
        info.cell(row=info.max_row, column=1).font = Font(bold=True)
    info["B2"].number_format = "DD-MM-YYYY HH:MM"
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 45

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer, total_rows
